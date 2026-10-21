"""
Generates the 5 BFS reference documents used by every lab in this series
(Nirikshan, Smriti, Drishti). Re-created here because this session does not
carry forward files from earlier labs - content is kept consistent with
those labs so the same reference agent and vector store behave identically.

Run once: python make_documents.py
"""
import json
import csv
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from docx import Document

OUT = Path(__file__).parent / "documents"
OUT.mkdir(exist_ok=True)


def make_pdf():
    path = OUT / "01_mortgage_lending_sop.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    styles = getSampleStyleSheet()
    body = [
        Paragraph("Lloyds Financial (Training Scenario) - Mortgage Lending SOP", styles["Title"]),
        Spacer(1, 12),
        Paragraph("1. Scope", styles["Heading2"]),
        Paragraph(
            "This SOP governs residential mortgage lending decisions for the Premier "
            "and Standard product lines. It is a training-scenario document and does "
            "not represent a real Lloyds policy.",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("2. Affordability Assessment", styles["Heading2"]),
        Paragraph(
            "Every application must pass a stressed-rate affordability check at the "
            "product rate plus 2.0 percentage points. The maximum loan-to-income "
            "multiple is 4.5x gross annual income for Standard, and 5.0x for Premier "
            "customers who hold an active Premier Current Account for 12+ months.",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("3. Fixed Rate Mortgage - Current Terms", styles["Heading2"]),
        Paragraph(
            "The 2-Year Fixed Rate Mortgage is offered at 4.79% APR for loan-to-value "
            "(LTV) up to 75%, and 5.09% APR for LTV between 75% and 90%. An early "
            "repayment charge of 2% of the outstanding balance applies if the mortgage "
            "is redeemed within the fixed period.",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("4. Overdraft Facility - Current Terms", styles["Heading2"]),
        Paragraph(
            "Arranged overdrafts on a Premier Current Account carry an Effective "
            "Annual Rate (EAR) of 19.9%. This rate is reviewed quarterly and was last "
            "confirmed current as of this document's publication.",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("5. Fraud Referral Threshold", styles["Heading2"]),
        Paragraph(
            "Any single transaction exceeding GBP 8,000, or any account showing three "
            "or more transactions above GBP 2,000 within a rolling 24-hour window, must "
            "be referred to the Fraud Operations queue before funds are released.",
            styles["BodyText"],
        ),
    ]
    doc.build(body)
    print("wrote", path)


def make_docx():
    path = OUT / "02_premier_account_terms.docx"
    doc = Document()
    doc.add_heading("Premier Current Account - Terms Summary (Training Scenario)", level=1)
    doc.add_heading("Eligibility", level=2)
    doc.add_paragraph(
        "Available to customers paying in at least GBP 75,000 per year, or holding "
        "GBP 100,000+ in linked savings and investments."
    )
    doc.add_heading("Monthly Fee", level=2)
    doc.add_paragraph("GBP 17 per month, waived if the eligibility criteria above are met.")
    doc.add_heading("Savings Products Linked to Premier", level=2)
    doc.add_paragraph(
        "Premier Easy Saver: 4.10% AER (variable), instant access.\n"
        "Premier Fixed Saver (1 year): 4.55% AER (fixed), no withdrawals until maturity."
    )
    doc.add_heading("Complaints Handling Commitment", level=2)
    doc.add_paragraph(
        "Premier complaints are acknowledged within 2 working days and target "
        "resolution within 10 working days. Unresolved complaints are escalated to "
        "a Premier Relationship Manager after 5 working days without progress."
    )
    doc.save(path)
    print("wrote", path)


def make_html():
    path = OUT / "03_customer_faq.html"
    html = """<!doctype html>
<html>
<head><title>Lloyds Financial - Customer FAQ (Training Scenario)</title></head>
<body>
<h1>Frequently Asked Questions</h1>

<h2>How do I contact support about my complaint?</h2>
<p>Email our complaints team at complaints-support@example-lloydsfinancial.co.uk
or call our customer line. For Premier customers, your Relationship Manager,
Priya Nair, can be reached directly on 0161 496 0147.</p>

<h2>What is the overdraft interest rate?</h2>
<p>The current arranged overdraft Effective Annual Rate (EAR) on a Premier
Current Account is 19.9%, reviewed quarterly.</p>

<h2>How long does a complaint take to resolve?</h2>
<p>Most complaints are resolved within 10 working days of acknowledgement. You
will receive a written update if it takes longer.</p>

<h2>Can I get a mortgage with a 10% deposit?</h2>
<p>Yes - our Standard and Premier mortgage products both lend up to 90% LTV,
subject to the stressed-rate affordability check described in the Mortgage
Lending SOP.</p>

<h2>What savings rate do I get with a Premier account?</h2>
<p>Premier Easy Saver pays 4.10% AER variable with instant access; the Premier
Fixed Saver pays 4.55% AER fixed for one year.</p>
</body>
</html>
"""
    path.write_text(html, encoding="utf-8")
    print("wrote", path)


def make_csv():
    path = OUT / "04_complaints_register.csv"
    rows = [
        ["complaint_id", "customer_ref", "category", "status", "opened", "summary"],
        ["CMP-1001", "CUST-5521", "Overdraft charge dispute", "open",
         "2026-09-28", "Customer disputes an overdraft interest charge applied after a payment date error."],
        ["CMP-1002", "CUST-5589", "Mortgage rate mis-quote", "open",
         "2026-10-01", "Customer was quoted the 75% LTV rate but offered the 90% LTV rate at completion."],
        ["CMP-1003", "CUST-5602", "Delayed fraud hold release", "open",
         "2026-10-03", "Legitimate transaction held for review beyond the standard turnaround; customer escalated."],
        ["CMP-1004", "CUST-5477", "Savings rate change notice", "resolved",
         "2026-09-10", "Customer queried notice period for a Fixed Saver rate change; resolved with explanation."],
        ["CMP-1005", "CUST-5390", "Premier fee charged in error", "resolved",
         "2026-08-30", "Monthly fee charged despite meeting waiver criteria; refunded same week."],
    ]
    with open(path, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)
    print("wrote", path)


def make_json():
    path = OUT / "05_savings_product_rates.json"
    data = {
        "as_of": "2026-10-01",
        "products": [
            {"name": "Premier Easy Saver", "aer": 4.10, "type": "variable", "access": "instant"},
            {"name": "Premier Fixed Saver (1yr)", "aer": 4.55, "type": "fixed", "access": "none_until_maturity"},
            {"name": "Standard Easy Saver", "aer": 3.25, "type": "variable", "access": "instant"},
            {"name": "ISA Saver", "aer": 3.90, "type": "variable", "access": "instant", "tax_wrapper": "ISA"},
        ],
    }
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")
    print("wrote", path)


if __name__ == "__main__":
    make_pdf()
    make_docx()
    make_html()
    make_csv()
    make_json()
